#!/usr/bin/env python3
"""Build the Excel implementation of Miniproject 2 (live formulas, Solver-ready).

Layout (the way the workbook would be built by hand in Excel):
    Data       WSJ quotes exactly as downloaded
    Bonds      one row per bond: 32nds conversion, coupon dates, accrued, YIELD, duration,
               model price, errors
    CashFlows  every remaining coupon/principal date of every bond and its present value
    Curve      Svensson parameters, Solver objective, fit statistics, zero-rate table, chart
    Notes      one-page methodology

    python build_excel.py --data ../data/Treasury_data_090426.xlsx \
        --results ../output/results.json --out ../excel/Miniproject2_ZCB_Term_Structure.xlsx
"""
from __future__ import annotations

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parent))
from zcb import load_wsj_quotes  # noqa: E402

LABOR_DAY_2026 = date(2026, 9, 7)

# --- layout constants shared with verify_excel.py ---------------------------------
DATA_HEADER_ROW = 6
DATA_FIRST_ROW = 7
BOND_HEADER_ROW = 7
BOND_FIRST_ROW = 8
CURVE_PARAM_FIRST_ROW = 4          # Curve!B4:B9 = beta_0 .. tau_2
CURVE_SSE_CELL = "B11"
CURVE_STATS = {"n_used": "B12", "price_rmse": "B13", "yield_rmse": "B14", "yield_mae": "B15", "yield_max": "B16"}
CURVE_TABLE_HEADER_ROW = 21
CURVE_TABLE_FIRST_ROW = 22
CF_DATE_FIRST_COL = 5              # CashFlows!E = date of coupon 1
BOND_COLS = {
    "maturity": "A", "coupon": "B", "bid": "C", "ask": "D", "wsj_yield": "E", "clean": "F", "prev": "G",
    "next": "H", "days_acc": "I", "days_period": "J", "accrued": "K", "dirty": "L", "ytm": "M", "ytm_diff": "N",
    "duration": "O", "n_coupons": "P", "years": "Q", "use": "R", "model_dirty": "S", "model_clean": "T",
    "price_err": "U", "model_ytm": "V", "yield_err": "W", "w_err": "X", "w_sq": "Y",
}
PARAM_NAMES = ["beta_0", "beta_1", "beta_2", "beta_3", "tau_1", "tau_2"]
PARAM_KEYS = ["beta0", "beta1", "beta2", "beta3", "tau1", "tau2"]

# --- plain Excel styling ------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
THIN = Side(style="thin", color="BFBFBF")
BOTTOM = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER_WRAP = Alignment(wrap_text=True, vertical="center", horizontal="center")
DATE_FMT = "m/d/yyyy"

ZERO_FORMULA = (
    "beta_0+beta_1*(1-EXP(-{t}/tau_1))/({t}/tau_1)"
    "+beta_2*((1-EXP(-{t}/tau_1))/({t}/tau_1)-EXP(-{t}/tau_1))"
    "+beta_3*((1-EXP(-{t}/tau_2))/({t}/tau_2)-EXP(-{t}/tau_2))"
)
FORWARD_FORMULA = (
    "beta_0+beta_1*EXP(-{t}/tau_1)+beta_2*({t}/tau_1)*EXP(-{t}/tau_1)+beta_3*({t}/tau_2)*EXP(-{t}/tau_2)"
)


def header(ws, row: int, labels: list[str], start_col: int = 1, height: float | None = None) -> None:
    for offset, label in enumerate(labels):
        cell = ws.cell(row=row, column=start_col + offset, value=label)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.border = BOTTOM
        cell.alignment = CENTER_WRAP
    if height:
        ws.row_dimensions[row].height = height


def widths(ws, spec: dict[str, float]) -> None:
    for col, width in spec.items():
        ws.column_dimensions[col].width = width


def name(wb: Workbook, label: str, ref: str) -> None:
    wb.defined_names[label] = DefinedName(label, attr_text=ref)


def as_dt(day: date | str) -> datetime:
    if isinstance(day, str):
        day = date.fromisoformat(day)
    return datetime.combine(day, datetime.min.time())


# ----------------------------------------------------------------------------- sheets
def build_data(wb: Workbook, quotes, quote_date: date) -> None:
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "U.S. Treasury Quotes"
    ws["A1"].font = TITLE
    ws["A2"] = quote_date.strftime("%A, %B %d, %Y")
    ws["A3"] = "Treasury Notes & Bonds (Wall Street Journal, source Tullett Prebon)"
    ws["A4"] = "Bid/Asked are in 32nds; the third decimal is eighths of a 32nd (e.g. 99.256 = 99 + 25 6/8 32nds)."
    header(ws, DATA_HEADER_ROW, ["Maturity", "Coupon", "Bid", "Asked", "Chg", "Asked yield"])
    for i, q in enumerate(quotes.itertuples(index=False)):
        r = DATA_FIRST_ROW + i
        ws.cell(row=r, column=1, value=as_dt(q.maturity)).number_format = DATE_FMT
        ws.cell(row=r, column=2, value=q.coupon).number_format = "0.000"
        ws.cell(row=r, column=3, value=q.bid_quote).number_format = "0.000"
        ws.cell(row=r, column=4, value=q.ask_quote).number_format = "0.000"
        ws.cell(row=r, column=5, value=q.chg)
        ws.cell(row=r, column=6, value=q.asked_yield).number_format = "0.000"
    ws.freeze_panes = f"A{DATA_FIRST_ROW}"
    widths(ws, {"A": 12, "B": 9, "C": 9, "D": 9, "E": 8, "F": 11})


def build_bonds(wb: Workbook, results: dict, n_bonds: int, n_cols: int) -> None:
    ws = wb.create_sheet("Bonds")
    ws["A1"] = "Bond-by-bond calculations"
    ws["A1"].font = TITLE
    ws["A2"] = "Quote date"
    ws["B2"] = as_dt(results["quote_date"])
    ws["A3"] = "Settlement date (T+1 business day)"
    ws["B3"] = "=WORKDAY(B2,1,B4)"
    ws["A4"] = "Market holiday (Labor Day)"
    ws["B4"] = as_dt(LABOR_DAY_2026)
    ws["A5"] = "Exclude bonds with less than this many years to maturity from the fit"
    ws["B5"] = results["svensson"]["min_years_in_fit"]
    for ref in ("B2", "B3", "B4", "B5"):
        ws[ref].number_format = DATE_FMT if ref != "B5" else "0.00"
    for ref in ("B2", "B4", "B5"):
        ws[ref].fill = INPUT_FILL
    name(wb, "Settle", "Bonds!$B$3")
    name(wb, "MinYears", "Bonds!$B$5")

    labels = [
        "Maturity", "Coupon (%)", "Bid (32nds)", "Asked (32nds)", "WSJ asked yield (%)", "Asked price (decimal)",
        "Previous coupon", "Next coupon", "Days accrued", "Days in period", "Accrued interest", "Dirty price",
        "YTM (%) from YIELD()", "YTM - WSJ (bp)", "Modified duration", "Coupons left", "Years to maturity",
        "Use in fit (1/0)", "Model dirty price", "Model clean price", "Price error", "Model YTM (%)",
        "Yield error (bp)", "Approx. yield error = price error / (dirty x duration)", "Squared, if used",
    ]
    header(ws, BOND_HEADER_ROW, labels, height=44)
    pv_first = get_column_letter(CF_DATE_FIRST_COL + n_cols + 1)
    pv_last = get_column_letter(CF_DATE_FIRST_COL + 2 * n_cols)
    C = BOND_COLS
    fmt = {
        "maturity": DATE_FMT, "coupon": "0.000", "bid": "0.000", "ask": "0.000", "wsj_yield": "0.000",
        "clean": "0.0000", "prev": DATE_FMT, "next": DATE_FMT, "accrued": "0.0000", "dirty": "0.0000",
        "ytm": "0.000", "ytm_diff": "0.00", "duration": "0.000", "years": "0.000", "model_dirty": "0.0000",
        "model_clean": "0.0000", "price_err": "0.0000", "model_ytm": "0.000", "yield_err": "0.00",
        "w_err": "0.00000", "w_sq": "0.0000000",
    }
    for i in range(n_bonds):
        r = BOND_FIRST_ROW + i
        d = DATA_FIRST_ROW + i
        formulas = {
            "maturity": f"=Data!A{d}",
            "coupon": f"=Data!B{d}",
            "bid": f"=Data!C{d}",
            "ask": f"=Data!D{d}",
            "wsj_yield": f"=Data!F{d}",
            # 32nds: thousandths digits "xxe" -> xx 32nds + e eighths of a 32nd (locale independent)
            "clean": f"=INT(D{r})+(INT(MOD(ROUND(D{r}*1000,0),1000)/10)+MOD(ROUND(D{r}*1000,0),10)/8)/32",
            "prev": f"=COUPPCD(Settle,A{r},2,1)",
            "next": f"=COUPNCD(Settle,A{r},2,1)",
            "days_acc": f"=Settle-G{r}",
            "days_period": f"=H{r}-G{r}",
            "accrued": f"=B{r}/2*I{r}/J{r}",
            "dirty": f"=F{r}+K{r}",
            "ytm": f"=YIELD(Settle,A{r},B{r}/100,F{r},100,2,1)*100",
            "ytm_diff": f"=(M{r}-E{r})*100",
            "duration": f"=MDURATION(Settle,A{r},B{r}/100,M{r}/100,2,1)",
            "n_coupons": f"=COUPNUM(Settle,A{r},2,1)",
            "years": f"=(A{r}-Settle)/365",
            "use": f"=IF(Q{r}>=MinYears,1,0)",
            "model_dirty": f"=SUM(CashFlows!{pv_first}{r}:{pv_last}{r})",
            "model_clean": f"=S{r}-K{r}",
            "price_err": f"=T{r}-F{r}",
            "model_ytm": f"=YIELD(Settle,A{r},B{r}/100,T{r},100,2,1)*100",
            "yield_err": f"=(V{r}-M{r})*100",
            "w_err": f"=(S{r}-L{r})/(L{r}*O{r})",
            "w_sq": f"=R{r}*X{r}^2",
        }
        for key, formula in formulas.items():
            cell = ws[f"{C[key]}{r}"]
            cell.value = formula
            if key in fmt:
                cell.number_format = fmt[key]
    last = BOND_FIRST_ROW + n_bonds - 1
    ws[f"W{last + 2}"] = "Sum of squares (objective):"
    ws[f"W{last + 2}"].font = BOLD
    ws[f"Y{last + 2}"] = f"=SUM(Y{BOND_FIRST_ROW}:Y{last})"
    ws[f"Y{last + 2}"].number_format = "0.000000"
    ws[f"Y{last + 2}"].font = BOLD
    ws.freeze_panes = f"C{BOND_FIRST_ROW}"
    widths(ws, {"A": 11, "B": 9, "C": 9, "D": 9, "E": 10, "F": 11, "G": 11, "H": 11, "I": 8, "J": 8, "K": 10,
                "L": 10, "M": 10, "N": 9, "O": 9, "P": 8, "Q": 9, "R": 8, "S": 11, "T": 11, "U": 10, "V": 10,
                "W": 9, "X": 10, "Y": 11})


def build_cashflows(wb: Workbook, n_bonds: int, n_cols: int) -> None:
    ws = wb.create_sheet("CashFlows")
    ws["A1"] = "Remaining cash flows of each bond and their present values off the Svensson curve"
    ws["A1"].font = TITLE
    ws["A2"] = ("Coupon k (k = 1..N) falls 6(k-1) months after the next coupon date; month-end bonds stay on month-ends. "
                "PV = (coupon/2 + 100 at maturity) x exp(-r(t) t), with t = (date - settlement)/365 and r(t) the Svensson zero rate "
                "using the parameters on the Curve sheet. Row sums feed Bonds!S (model dirty price).")
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:P2")
    ws.row_dimensions[2].height = 45
    header(ws, BOND_HEADER_ROW, ["Maturity", "Coupon (%)", "Coupons left (N)", "Month-end? (1/0)"], height=32)
    date0 = CF_DATE_FIRST_COL
    pv0 = CF_DATE_FIRST_COL + n_cols + 1
    header(ws, BOND_HEADER_ROW, [f"Date {k}" for k in range(1, n_cols + 1)], date0)
    header(ws, BOND_HEADER_ROW, [f"PV {k}" for k in range(1, n_cols + 1)], pv0)
    ws.cell(row=BOND_HEADER_ROW - 1, column=date0, value="Payment dates").font = BOLD
    ws.cell(row=BOND_HEADER_ROW - 1, column=pv0, value="Present values (per 100 face)").font = BOLD
    for i in range(n_bonds):
        r = BOND_FIRST_ROW + i
        ws[f"A{r}"] = f"=Bonds!A{r}"
        ws[f"A{r}"].number_format = DATE_FMT
        ws[f"B{r}"] = f"=Bonds!B{r}"
        ws[f"B{r}"].number_format = "0.000"
        ws[f"C{r}"] = f"=Bonds!P{r}"
        ws[f"D{r}"] = f"=IF(DAY(A{r})=DAY(EOMONTH(A{r},0)),1,0)"
        for k in range(1, n_cols + 1):
            dcol = get_column_letter(date0 + k - 1)
            pcol = get_column_letter(pv0 + k - 1)
            shifted = f"EDATE(Bonds!$H{r},6*({k}-1))"
            ws[f"{dcol}{r}"] = f'=IF({k}>$C{r},"",IF($D{r}=1,EOMONTH({shifted},0),{shifted}))'
            ws[f"{dcol}{r}"].number_format = DATE_FMT
            t = f"(({dcol}{r}-Settle)/365)"
            amount = f"($B{r}/2+IF({k}=$C{r},100,0))"
            ws[f"{pcol}{r}"] = f'=IF({dcol}{r}="",0,{amount}*EXP(-({ZERO_FORMULA.format(t=t)})*{t}))'
            ws[f"{pcol}{r}"].number_format = "0.0000"
    ws.freeze_panes = f"E{BOND_FIRST_ROW}"
    widths(ws, {"A": 11, "B": 9, "C": 8, "D": 9})
    for c in range(date0, pv0 + n_cols):
        ws.column_dimensions[get_column_letter(c)].width = 10


def build_curve(wb: Workbook, results: dict, n_bonds: int) -> None:
    ws = wb.create_sheet("Curve")
    ws["A1"] = "Svensson zero-coupon curve: parameters, fit and results"
    ws["A1"].font = TITLE
    ws["A3"] = "Parameters (continuously compounded, decimals). Run Solver on B11 or edit by hand."
    ws["A3"].font = BOLD
    params = results["svensson"]["params"]
    notes = ["level: r(t) as t -> infinity", "slope: r(0) = beta_0 + beta_1", "first hump (decay tau_1)",
             "second hump (decay tau_2)", "decay of slope / first hump, years", "decay of second hump, years"]
    for i, (label, key, note) in enumerate(zip(PARAM_NAMES, PARAM_KEYS, notes)):
        r = CURVE_PARAM_FIRST_ROW + i
        ws[f"A{r}"] = label
        ws[f"B{r}"] = float(params[key])
        ws[f"B{r}"].number_format = "0.000000"
        ws[f"B{r}"].fill = INPUT_FILL
        ws[f"C{r}"] = note
        name(wb, label, f"Curve!$B${r}")

    first, last = BOND_FIRST_ROW, BOND_FIRST_ROW + n_bonds - 1
    R, U, W = f"Bonds!R{first}:R{last}", f"Bonds!U{first}:U{last}", f"Bonds!W{first}:W{last}"
    stats = [
        (11, "Objective: sum over used bonds of ((model - market dirty price) / (dirty price x duration))^2 = sum of squared first-order yield errors", f"=Bonds!Y{last + 2}", "0.00000000"),
        (12, "Bonds used in the fit", f"=SUM({R})", "0"),
        (13, "Price RMSE (per 100 face)", f"=SQRT(SUMPRODUCT({R},{U},{U})/B12)", "0.0000"),
        (14, "Yield RMSE (bp)", f"=SQRT(SUMPRODUCT({R},{W},{W})/B12)", "0.00"),
        (15, "Yield mean absolute error (bp)", f"=SUMPRODUCT({R},ABS({W}))/B12", "0.00"),
        (16, "Largest absolute yield error (bp)", f"=MAX(INDEX({R}*ABS({W}),0))", "0.00"),
        (17, "r(0) = beta_0 + beta_1  (%)", "=(beta_0+beta_1)*100", "0.000"),
        (18, "r(30 years)  (%)", "=(" + ZERO_FORMULA.format(t="30") + ")*100", "0.000"),
    ]
    for r, label, formula, fmt in stats:
        ws[f"A{r}"] = label
        ws[f"B{r}"] = formula
        ws[f"B{r}"].number_format = fmt
    ws["B11"].font = BOLD
    ws["D11"] = ("Solver: Set Objective B11 to Min; By Changing B4:B9; Constraints B4 >= 0, B8 >= 0.05, B9 >= 0.05; "
                 "GRG Nonlinear; untick 'Make Unconstrained Variables Non-Negative'. Try a few starting values for "
                 "tau_1 (1-4) and tau_2 (6-20) and keep the lowest B11.")
    ws["D11"].alignment = WRAP
    ws.merge_cells("D11:G14")

    ws[f"A{CURVE_TABLE_HEADER_ROW - 1}"] = ("Zero rate, discount factor and instantaneous forward rate at every Treasury payment date in the sample "
                                             "(dates: the distinct dates on the CashFlows sheet, listed as values)")
    ws[f"A{CURVE_TABLE_HEADER_ROW - 1}"].font = BOLD
    header(ws, CURVE_TABLE_HEADER_ROW, ["Payment date", "Days", "t (years)", "Zero rate r(t) %", "Discount factor", "Forward f(t) %"], height=32)
    pay = results["payment_dates"]
    for i, p in enumerate(pay):
        r = CURVE_TABLE_FIRST_ROW + i
        ws[f"A{r}"] = as_dt(p["payment_date"])
        ws[f"A{r}"].number_format = DATE_FMT
        ws[f"B{r}"] = f"=A{r}-Settle"
        ws[f"C{r}"] = f"=B{r}/365"
        ws[f"C{r}"].number_format = "0.0000"
        ws[f"D{r}"] = "=(" + ZERO_FORMULA.format(t=f"C{r}") + ")*100"
        ws[f"D{r}"].number_format = "0.0000"
        ws[f"E{r}"] = f"=EXP(-D{r}/100*C{r})"
        ws[f"E{r}"].number_format = "0.000000"
        ws[f"F{r}"] = "=(" + FORWARD_FORMULA.format(t=f"C{r}") + ")*100"
        ws[f"F{r}"].number_format = "0.0000"
    table_last = CURVE_TABLE_FIRST_ROW + len(pay) - 1
    widths(ws, {"A": 13, "B": 12, "C": 10, "D": 13, "E": 13, "F": 12, "G": 12})

    chart = ScatterChart()
    chart.title = "Treasury zero-coupon curve (continuous) vs. bond yields to maturity"
    chart.style = 2
    chart.x_axis.title = "Years from settlement"
    chart.y_axis.title = "Rate (%)"
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = 0, 31
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = 2.5, 6.0
    chart.x_axis.majorUnit = 5
    chart.height, chart.width = 10, 20
    x_ref = Reference(ws, min_col=3, min_row=CURVE_TABLE_FIRST_ROW, max_row=table_last)
    zero = Series(Reference(ws, min_col=4, min_row=CURVE_TABLE_FIRST_ROW, max_row=table_last), x_ref, title="Zero rate r(t)")
    zero.marker.symbol = "none"
    zero.graphicalProperties.line.width = 22000
    fwd = Series(Reference(ws, min_col=6, min_row=CURVE_TABLE_FIRST_ROW, max_row=table_last), x_ref, title="Forward rate f(t)")
    fwd.marker.symbol = "none"
    fwd.graphicalProperties.line.dashStyle = "dash"
    fwd.graphicalProperties.line.width = 15000
    bonds_ws = wb["Bonds"]
    ytm = Series(Reference(bonds_ws, min_col=13, min_row=first, max_row=last),
                 Reference(bonds_ws, min_col=17, min_row=first, max_row=last), title="Bond YTM (street)")
    ytm.marker.symbol = "circle"
    ytm.marker.size = 4
    ytm.graphicalProperties.line.noFill = True
    for s in (zero, fwd, ytm):
        chart.series.append(s)
    ws.add_chart(chart, "H20")
    ws.freeze_panes = f"A{CURVE_TABLE_FIRST_ROW}"


def notes_text(results: dict) -> list[str]:
    sv = results["svensson"]
    p = sv["params"]
    pay = results["payment_dates"]
    conv = results["conventions"][0]
    n_total, n_used = sv["n_bonds_total"], sv["n_bonds_in_fit"]

    def zero_at(years: float) -> float:
        return min(pay, key=lambda r: abs(r["t_years"] - years))["zero_rate_cc_pct"]

    prof = {round(r["beta0_fixed_pct"]): r for r in results["beta0_profile"]}
    shift = max(abs(prof[3][k] - prof[0][k]) for k in prof[0] if k.startswith("zero_")) * 100
    others = [c["rmse_bp"] for c in results["conventions"]
              if c["price_format"] == conv["price_format"] and c["settlement"] != conv["settlement"]]
    alt_lo, alt_hi = min(others), max(others)
    dec_lo = min(c["rmse_bp"] for c in results["conventions"] if c["price_format"] == "plain decimal")
    return [
        f"Goal. Build a continuously compounded zero-coupon (ZCB) term structure from the WSJ Treasury note and bond quotes "
        f"of {date.fromisoformat(results['quote_date']):%A, %B %d, %Y} and report the discount rate for every coupon and "
        f"principal payment date in the sample.",
        f"Data (Data sheet). {n_total} notes and bonds. Prices are quoted in 32nds and the third decimal is eighths of a 32nd, "
        f"so 99.256 = 99 + 25.75/32 (converted in Bonds!F). We price off the asked side. Settlement is the next business day, "
        f"{date.fromisoformat(results['settlement']):%B %d, %Y} (Monday {LABOR_DAY_2026:%B %d} is Labor Day). As a check, "
        f"recomputing each bond's yield with YIELD() reproduces the WSJ asked yield with a median error of "
        f"{conv['median_abs_bp']:.2f} bp (Bonds!N); other settlement dates give {alt_lo:.1f} to {alt_hi:.1f} bp and reading the quotes as "
        f"plain decimals gives over {dec_lo:.0f} bp.",
        "Cash flows (Bonds, CashFlows sheets). Coupons are semi-annual on the maturity day of month (month-end maturities pay "
        "on month-ends) and principal of 100 is paid at maturity. Accrued interest = coupon/2 x days since the last coupon / "
        "days in the coupon period (actual/actual, COUPPCD and COUPNCD). Dirty price = clean price + accrued. Time to each "
        "payment is t = days/365 from settlement.",
        "Model. The zero rate follows the Svensson (1994) form suggested with the data: r(t) = beta_0 + beta_1 (1-e^(-t/tau_1))/(t/tau_1) "
        "+ beta_2 [(1-e^(-t/tau_1))/(t/tau_1) - e^(-t/tau_1)] + beta_3 [(1-e^(-t/tau_2))/(t/tau_2) - e^(-t/tau_2)]. The discount "
        "factor is d(t) = e^(-r(t) t), and each bond's model price is the sum of its payments times their discount factors "
        "(CashFlows sheet, one column per remaining coupon; row sums feed Bonds!S). r(0) = beta_0 + beta_1, r(infinity) = beta_0, "
        "and the instantaneous forward rate is f(t) = beta_0 + beta_1 e^(-t/tau_1) + beta_2 (t/tau_1) e^(-t/tau_1) + beta_3 (t/tau_2) e^(-t/tau_2).",
        f"Fitting (Curve sheet). Solver minimises the sum over bonds of ((model dirty price - market dirty price) / (dirty price x modified duration))^2 "
        f"over the six parameters, with beta_0 >= 0 and tau_1, tau_2 >= 0.05. Each price error divided by price times duration is the "
        f"first-order yield error, so this is the sum of squared yield errors and a 30-year bond and a 1-year note count equally. The {n_total - n_used} bonds with under 3 months "
        f"to maturity are left out of the objective (their duration is close to zero, so a one-cent quote error is a 50 bp yield "
        f"error) but are still priced off the curve as an out-of-sample check. Because the objective has several local minima in "
        f"(tau_1, tau_2), Solver was run from several starting values for the two decay parameters and the lowest objective kept.",
        f"Results. beta_0 = {p['beta0']:.4f} (the non-negativity constraint binds), beta_1 = {p['beta1']:.4f}, beta_2 = {p['beta2']:.4f}, "
        f"beta_3 = {p['beta3']:.4f}, tau_1 = {p['tau1']:.2f}, tau_2 = {p['tau2']:.2f}. The curve prices the {n_used} bonds with a yield "
        f"RMSE of {sv['yield_rmse_bp']:.1f} bp (mean absolute {sv['yield_mae_bp']:.1f} bp, largest {sv['yield_max_abs_bp']:.1f} bp). "
        f"The zero curve rises from {(p['beta0'] + p['beta1']) * 100:.2f}% at the short end to {zero_at(1):.2f}% at 1 year, "
        f"{zero_at(5):.2f}% at 5, {zero_at(10):.2f}% at 10, {zero_at(20):.2f}% at 20 and {zero_at(30):.2f}% at 30 years. The "
        f"Curve sheet lists r(t), d(t) and f(t) for all {len(pay)} distinct payment dates from "
        f"{date.fromisoformat(pay[0]['payment_date']):%b %d, %Y} to {date.fromisoformat(pay[-1]['payment_date']):%b %d, %Y}, with the chart. "
        f"One caveat: beta_0 (the rate at infinite maturity) is not well identified by 30 years of data; fixing it anywhere between 0% "
        f"and 3% worsens the fit by only about {prof[3]['approx_yield_rmse_bp'] - prof[0]['approx_yield_rmse_bp']:.1f} bp RMSE and moves "
        f"the curve inside the sample by about {shift:.0f} bp at the two ends, so the payment-date discount factors are robust to it. "
        f"The curve should not be extrapolated beyond the 30-year sample.",
    ]


def build_notes(wb: Workbook, results: dict) -> None:
    ws = wb.create_sheet("Notes")
    ws["A1"] = "Miniproject 2 - Creating a ZCB Term Structure"
    ws["A1"].font = TITLE
    ws["A2"] = "FRE 6103 Valuation for Financial Engineering. Methodology and results."
    for i, paragraph in enumerate(notes_text(results)):
        r = 4 + i
        cell = ws.cell(row=r, column=1, value=paragraph)
        cell.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        ws.row_dimensions[r].height = 15 * (len(paragraph) // 105 + 2)
    for col in "ABCDEFGHIJ":
        ws.column_dimensions[col].width = 14


def build_workbook(data_path: Path, results_path: Path, out_path: Path) -> Path:
    results = json.loads(results_path.read_text())
    sheet = load_wsj_quotes(data_path)
    n_bonds = len(sheet.quotes)
    n_cols = max(b["n_cash_flows"] for b in results["bonds"])
    wb = Workbook()
    build_data(wb, sheet.quotes, sheet.quote_date)
    build_bonds(wb, results, n_bonds, n_cols)
    build_cashflows(wb, n_bonds, n_cols)
    build_curve(wb, results, n_bonds)
    build_notes(wb, results)
    wb.calculation.fullCalcOnLoad = True
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--data", default="../data/Treasury_data_090426.xlsx")
    parser.add_argument("--results", default="../output/results.json")
    parser.add_argument("--out", default="../excel/Miniproject2_ZCB_Term_Structure.xlsx")
    args = parser.parse_args(argv)
    path = build_workbook(Path(args.data), Path(args.results), Path(args.out))
    print(f"Wrote {path.resolve()} ({path.stat().st_size / 1e6:.1f} MB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
